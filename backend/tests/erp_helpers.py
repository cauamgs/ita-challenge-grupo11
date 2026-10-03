"""Gera variações da planilha do ERP para simular mudanças entre sincronizações."""

import io

from openpyxl import load_workbook

from tests.conftest import SEED_DIR

XLSX = SEED_DIR / "base_popular_pet.xlsx"
HEADER_ROW = 3  # 2 linhas de título antes do cabeçalho


def modified_workbook(changes: dict[str, list[tuple[dict, dict]]]) -> bytes:
    """changes = {aba: [(filtro {coluna: valor}, novos valores {coluna: valor}), ...]}"""
    book = load_workbook(XLSX)
    for sheet_name, edits in changes.items():
        sheet = book[sheet_name]
        header = {cell.value: cell.column for cell in sheet[HEADER_ROW] if cell.value}
        for where, values in edits:
            for row in sheet.iter_rows(min_row=HEADER_ROW + 1):
                if all(row[header[k] - 1].value == v for k, v in where.items()):
                    for k, v in values.items():
                        row[header[k] - 1].value = v
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def without_sheet(sheet_name: str) -> bytes:
    book = load_workbook(XLSX)
    del book[sheet_name]
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()
